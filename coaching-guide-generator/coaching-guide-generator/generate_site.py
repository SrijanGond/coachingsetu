#!/usr/bin/env python3
"""
Generate the whole pan-India coaching guide site from one spreadsheet.

Usage:
    pip install openpyxl --break-system-packages
    python generate_site.py institutes_master.xlsx output/

Reads the "Institutes" sheet and writes:
    output/index.html                              (state/city picker)
    output/style.css
    output/<state-slug>/<city-slug>/index.html
    output/<state-slug>/<city-slug>/ssc-coaching-<city-slug>.html
    output/<state-slug>/<city-slug>/bank-po-coaching-<city-slug>.html
    output/<state-slug>/<city-slug>/upsc-coaching-<city-slug>.html
"""
import sys
import re
import shutil
from pathlib import Path
from collections import defaultdict
import openpyxl

BRAND = "CoachingSetu"
EXAM_LABELS = {"SSC": "SSC", "BankPO": "Bank PO", "UPSC": "UPSC/PCS"}
EXAM_SLUGS = {"SSC": "ssc-coaching", "BankPO": "bank-po-coaching", "UPSC": "upsc-coaching"}


def slugify(text):
    text = text.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


def fmt_fee(val):
    if val in (None, ""):
        return "Fee to be confirmed"
    try:
        return f"\u20b9{int(val):,}"
    except (ValueError, TypeError):
        return str(val)


def load_institutes(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Institutes"]
    rows = []
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    def get(row, col_name):
        i = idx.get(col_name)
        return row[i] if i is not None else None

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["State"]] or not row[idx["City"]] or not row[idx["Institute Name"]]:
            continue
        exams = [e.strip() for e in (get(row, "Exams (comma list: SSC,BankPO,UPSC)") or "").split(",") if e.strip()]
        featured_on = [e.strip() for e in (get(row, "Featured On (comma list: SSC,BankPO,UPSC or blank)") or "").split(",") if e.strip()]
        rows.append({
            "state": get(row, "State").strip(),
            "city": get(row, "City").strip(),
            "name": get(row, "Institute Name").strip(),
            "exams": exams,
            "coverage": {
                "SSC": get(row, "SSC Coverage (e.g. CGL, CHSL, MTS)") or "",
                "BankPO": get(row, "BankPO Coverage") or "",
                "UPSC": get(row, "UPSC Coverage") or "",
            },
            "timing": get(row, "Batch Timing") or "Timing to be confirmed",
            "fee": {
                "SSC": get(row, "Fee SSC (Rs)"),
                "BankPO": get(row, "Fee BankPO (Rs)"),
                "UPSC": get(row, "Fee UPSC (Rs)"),
            },
            "years": get(row, "Years in Operation") or "",
            "website": get(row, "Website / Listing") or "",
            "featured_on": featured_on,
            "notes": get(row, "Notes") or "",
        })
    return rows


PAGE_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{description}">
<link rel="stylesheet" href="/style.css">
</head>
<body>

<header class="site">
  <div class="wrap">
    <a class="brand" href="/index.html">{brand}</a>
    <nav class="site">
{nav}
    </nav>
  </div>
</header>

<section class="hero">
  <div class="wrap">
    <h1>{h1}</h1>
    <p class="lede">{lede}</p>
    <a class="btn" href="#enquire">Get Free Counselling</a>
  </div>
</section>

<section>
  <div class="wrap">
    <h2>{table_heading}</h2>
    <table class="compare">
      <thead>
        <tr>{table_headers}</tr>
      </thead>
      <tbody id="institute-table-body">
{table_rows}
      </tbody>
    </table>
    {table_footnote}
  </div>
</section>

<section id="enquire">
  <div class="wrap">
    <h2>Get free guidance</h2>
    <form class="lead">
      <label for="name">Full name</label>
      <input id="name" name="name" type="text" required>
      <label for="phone">Phone number</label>
      <input id="phone" name="phone" type="tel" required>
      <label for="batch">Preferred batch time</label>
      <select id="batch" name="batch">
        <option>Morning</option>
        <option>Evening</option>
        <option>Weekend</option>
      </select>
      <input type="hidden" name="exam" value="{exam_value}">
      <input type="hidden" id="routed-to" name="routed_to" value="">
      <button class="btn" type="submit">Submit — get a call back</button>
    </form>
  </div>
</section>

<footer class="site">
  <div class="wrap">
{footer_links}
    <div style="margin-top:8px;">&copy; 2026 {brand}</div>
  </div>
</footer>

<script>
  const FEATURED_INSTITUTE = "{featured_slug}"; // set by generate_site.py from the spreadsheet
  (function applyFeatured() {{
    const params = new URLSearchParams(window.location.search);
    const featured = params.get("featured") || FEATURED_INSTITUTE;
    if (!featured) return;
    const row = document.querySelector(`#institute-table-body tr[data-institute="${{featured}}"]`);
    if (!row) return;
    row.classList.add("featured-row");
    const firstCell = row.querySelector("td");
    if (firstCell && !firstCell.querySelector(".badge")) {{
      const badge = document.createElement("span");
      badge.className = "badge";
      badge.textContent = "Featured";
      firstCell.appendChild(badge);
    }}
    row.parentNode.prepend(row);
    const routedInput = document.getElementById("routed-to");
    if (routedInput) routedInput.value = featured;
  }})();
</script>

</body>
</html>
"""


def build_nav(city_slug_path, active=None):
    items = [
        (f"/{city_slug_path}/index.html", "Home"),
        (f"/{city_slug_path}/ssc-coaching-{city_slug_path.split('/')[-1]}.html", "SSC Coaching"),
        (f"/{city_slug_path}/bank-po-coaching-{city_slug_path.split('/')[-1]}.html", "Bank PO"),
        (f"/{city_slug_path}/upsc-coaching-{city_slug_path.split('/')[-1]}.html", "UPSC"),
        ("#enquire", "Free Counselling"),
    ]
    return "\n".join(f'      <a href="{href}">{label}</a>' for href, label in items)


def build_footer(city_slug_path, city_name):
    slug = city_slug_path.split("/")[-1]
    items = [
        (f"/{city_slug_path}/ssc-coaching-{slug}.html", f"SSC Coaching {city_name}"),
        (f"/{city_slug_path}/bank-po-coaching-{slug}.html", f"Bank PO Coaching {city_name}"),
        (f"/{city_slug_path}/upsc-coaching-{slug}.html", f"UPSC Coaching {city_name}"),
        ("#enquire", "Contact"),
    ]
    return "\n".join(f'    <a href="{href}">{label}</a>' for href, label in items)


def render_table_rows(institutes, exam):
    rows_html = []
    for inst in institutes:
        if exam and exam not in inst["exams"]:
            continue
        slug = slugify(inst["name"])
        coverage = inst["coverage"].get(exam, "") if exam else ", ".join(
            e for e in [inst["coverage"].get(x, "") for x in inst["exams"]] if e
        )
        fee = fmt_fee(inst["fee"].get(exam)) if exam else " / ".join(
            fmt_fee(inst["fee"].get(x)) for x in inst["exams"] if inst["fee"].get(x)
        ) or "Fee to be confirmed"
        rows_html.append(
            f'        <tr data-institute="{slug}"><td>{inst["name"]}</td>'
            f'<td>{coverage or "General"}</td><td>{inst["timing"]}</td><td>{fee}</td></tr>'
        )
    return "\n".join(rows_html) if rows_html else '        <tr><td colspan="4">No institutes added yet — add rows to the spreadsheet.</td></tr>'


def find_featured(institutes, exam):
    for inst in institutes:
        if exam in inst.get("featured_on", []):
            return slugify(inst["name"])
    return ""


def write_page(path, **kwargs):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(PAGE_TEMPLATE.format(**kwargs), encoding="utf-8")


def generate(xlsx_path, out_dir):
    out_dir = Path(out_dir)
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True)

    institutes = load_institutes(xlsx_path)

    # copy shared stylesheet
    style_src = Path(__file__).parent / "style.css"
    shutil.copy(style_src, out_dir / "style.css")

    by_state_city = defaultdict(list)
    for inst in institutes:
        by_state_city[(inst["state"], inst["city"])].append(inst)

    # --- per-city pages ---
    for (state, city), city_institutes in by_state_city.items():
        state_slug = slugify(state)
        city_slug = slugify(city)
        city_path = f"{state_slug}/{city_slug}"

        # Homepage
        write_page(
            out_dir / city_path / "index.html",
            title=f"Best SSC, Bank PO & UPSC Coaching Institutes in {city} ({state}) - {BRAND}",
            description=f"Compare top coaching institutes in {city}, {state} for SSC, Banking, and UPSC exams.",
            brand=BRAND,
            nav=build_nav(city_path),
            h1=f"Find the right coaching institute in {city} before you enrol",
            lede=f"We compare batch timings, fees, and specialisations across {city}'s leading SSC, Bank PO, and UPSC institutes.",
            table_heading=f"Top coaching institutes in {city}",
            table_headers="<th>Institute</th><th>Coverage</th><th>Batch timing</th><th>Fee (approx.)</th>",
            table_rows=render_table_rows(city_institutes, None),
            table_footnote='<p style="font-size:.85rem;color:#777;margin-top:10px;">Fees marked "to be confirmed" need a follow-up call.</p>',
            exam_value="General",
            featured_slug="",
            footer_links=build_footer(city_path, city),
        )

        # Exam pages
        for exam, label in EXAM_LABELS.items():
            slug_name = EXAM_SLUGS[exam]
            write_page(
                out_dir / city_path / f"{slug_name}-{city_slug}.html",
                title=f"Best {label} Coaching in {city} - Compare Institutes & Fees",
                description=f"Find the best {label} coaching institutes in {city}. Compare batch timings, fees and get free counselling.",
                brand=BRAND,
                nav=build_nav(city_path),
                h1=f"{label} coaching in {city} — compare top institutes",
                lede=f"Comparing {label} coaching institutes in {city} — batch timings, fees, and what to look for.",
                table_heading=f"{label} coaching institutes in {city}",
                table_headers="<th>Institute</th><th>Coverage</th><th>Batch timing</th><th>Fee (approx.)</th>",
                table_rows=render_table_rows(city_institutes, exam),
                table_footnote='<p style="font-size:.85rem;color:#777;margin-top:10px;">Fees marked "to be confirmed" need a follow-up call.</p>',
                exam_value=label,
                featured_slug=find_featured(city_institutes, exam),
                footer_links=build_footer(city_path, city),
            )

    # --- root state/city picker ---
    states = defaultdict(list)
    for (state, city) in by_state_city:
        states[state].append(city)

    rows = []
    for state in sorted(states):
        state_slug = slugify(state)
        for city in sorted(states[state]):
            city_slug = slugify(city)
            rows.append(
                f'<tr><td>{state}</td><td>{city}</td>'
                f'<td><a href="/{state_slug}/{city_slug}/index.html">View {city} →</a></td></tr>'
            )
    root_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{BRAND} — Compare Coaching Institutes Across India</title>
<meta name="description" content="Compare SSC, Bank PO, and UPSC coaching institutes state by state and city by city across India.">
<link rel="stylesheet" href="/style.css">
</head>
<body>
<header class="site"><div class="wrap"><a class="brand" href="/index.html">{BRAND}</a></div></header>
<section class="hero"><div class="wrap">
  <h1>Compare coaching institutes across India</h1>
  <p class="lede">Pick your state and city to compare SSC, Bank PO, and UPSC coaching institutes.</p>
</div></section>
<section><div class="wrap">
  <h2>Choose your city</h2>
  <table class="compare"><thead><tr><th>State</th><th>City</th><th>Explore</th></tr></thead>
  <tbody>
{chr(10).join('    ' + r for r in rows)}
  </tbody></table>
</div></section>
<footer class="site"><div class="wrap"><div>&copy; 2026 {BRAND}</div></div></footer>
</body>
</html>
"""
    (out_dir / "index.html").write_text(root_html, encoding="utf-8")
    print(f"Generated {len(by_state_city)} cities across {len(states)} state(s) into {out_dir}/")


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "institutes_master.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "output"
    generate(xlsx, out)
